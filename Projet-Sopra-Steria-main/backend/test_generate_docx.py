from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re
import shutil
import pytest
from docx import Document
from PyPDF2 import PdfReader
from analyser_cv import iter_document_text
from generators.generate_sopra_docx import generate_sopra_docx
from generators.pdf_sopra_profile import generate_sopra_profile_pdf
from generators.docx_to_pdf import convert_docx_to_pdf, ConversionUnavailable


@pytest.mark.parametrize("experience_count", [0, 1, 4])
def test_word_has_no_residual_placeholders_and_keeps_all_experiences(sample_cv, experience_count, monkeypatch, tmp_path):
    data = deepcopy(sample_cv)
    data["experiences"] = [{**sample_cv["experiences"][0], "entreprise": f"Entreprise exemple {i}"} for i in range(experience_count)]
    monkeypatch.chdir(tmp_path)  # Le modèle ne dépend pas du répertoire courant.
    buffer = BytesIO()
    generate_sopra_docx(data, buffer)
    buffer.seek(0)
    document = Document(buffer)
    text = "\n".join(iter_document_text(document))
    assert not re.search(r"\{\{.*?\}\}", text)
    assert data["contact"]["email"] in text
    for exp in data["experiences"]:
        assert exp["entreprise"] in text
    assert not document.core_properties.author
    assert not document.core_properties.last_modified_by


def test_pdf_wraps_long_content_and_keeps_last_experience(sample_cv):
    data = deepcopy(sample_cv)
    data["experiences"] = [{**sample_cv["experiences"][0], "description": ("Texte long <>& à conserver. " * 90) + f"FIN_MISSION_{i}"} for i in range(7)]
    buffer = BytesIO()
    generate_sopra_profile_pdf(data, buffer)
    buffer.seek(0)
    document = PdfReader(buffer)
    text = "\n".join(page.extract_text() or "" for page in document.pages)
    assert len(document.pages) > 2
    for i in range(7):
        assert f"FIN_MISSION_{i}" in text


def test_word_without_name_keeps_email():
    buffer = BytesIO()
    generate_sopra_docx({"contact": {"email": "demo@example.invalid"}}, buffer)
    buffer.seek(0)
    text = "\n".join(iter_document_text(Document(buffer)))
    assert "demo@example.invalid" in text
    assert "Environnement technique" not in text
    assert "{{" not in text


def test_converter_fails_explicitly_without_binary(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    Document().save(source)
    monkeypatch.setattr(shutil, "which", lambda _: None)
    monkeypatch.delenv("PROGRAMFILES", raising=False)
    monkeypatch.delenv("PROGRAMFILES(X86)", raising=False)
    with pytest.raises(ConversionUnavailable):
        convert_docx_to_pdf(source, tmp_path / "result.pdf")
    assert not (tmp_path / "result.pdf").exists()


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice non installé")
def test_libreoffice_conversion_uses_requested_output_name(tmp_path, sample_docx):
    source = tmp_path / "source.docx"
    source.write_bytes(sample_docx)
    target = tmp_path / "renamed.pdf"
    assert Path(convert_docx_to_pdf(source, target)) == target
    assert PdfReader(target).pages
