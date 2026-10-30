from copy import deepcopy
from io import BytesIO
from docx import Document
import pytest
from extractors.version_mapper import normalize_old_cv_to_new, convert_v2_to_old_format


def test_extracted_cv_is_preserved_and_idempotent(sample_cv):
    original = deepcopy(sample_cv)
    normalized = normalize_old_cv_to_new(sample_cv)
    assert normalized["contact"] == sample_cv["contact"]
    assert normalized["experiences"] == sample_cv["experiences"]
    assert normalized["formations"] == sample_cv["formations"]
    assert normalized["competences"] == sample_cv["competences"]
    assert normalized == normalize_old_cv_to_new(normalized)
    assert sample_cv == original
    exported = convert_v2_to_old_format(normalized)
    assert exported["contact"] == original["contact"]
    assert exported["experiences"] == original["experiences"]


def test_legacy_blocks_keep_missions_environment_and_contact():
    source = {"resultats_spacy": {"contact": {"nom": "Camille EXEMPLE", "email": "demo@example.invalid"},
              "competences_techniques": ["Python"], "experiences": [
                  "Janvier 2020 - Décembre 2022 - Entreprise Exemple", "Mission à préserver", "Environnement technique : Python, SQL"]}}
    normalized = normalize_old_cv_to_new(source)
    assert normalize_old_cv_to_new(normalized) == normalized
    exported = convert_v2_to_old_format(normalized)
    assert exported["experiences"][0]["description"] == "Mission à préserver"
    assert exported["experiences"][0]["environnement_technique"] == "Python, SQL"
    assert exported["contact"]["nom"] == "Camille EXEMPLE"


def test_undated_experience_is_not_silently_discarded():
    cv = normalize_old_cv_to_new({"experiences": ["Une expérience sans date", "Une mission"]})
    assert cv["experiences"][0]["titre"] == "Une expérience sans date"
    assert cv["experiences"][0]["missions"] == ["Une mission"]


def test_docx_normalization_reads_table_sections():
    doc = Document()
    doc.add_paragraph("Camille EXEMPLE")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Compétences techniques"
    table.cell(1, 0).text = "Python"
    data = BytesIO()
    doc.save(data)
    data.seek(0)
    assert normalize_old_cv_to_new({}, docx_path=data)["competences_techniques"] == ["Python"]


@pytest.mark.parametrize("data", [[], None, {"experiences": [42]}, {"competences": "Python"},
                                  {"formations": {}}, {"meta": []}, {"competences": [{}]},
                                  {"formations": [42]}, {"contact": {"nom": []}},
                                  {"experiences": [{"missions": "Une mission"}]}])
def test_invalid_schema_is_rejected(data):
    with pytest.raises(ValueError):
        normalize_old_cv_to_new(data)
