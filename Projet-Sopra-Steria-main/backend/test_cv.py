from io import BytesIO
from docx import Document
from analyser_cv import lire_cv_docx, extraire_infos_cv
from extractors.section_classifier import parse_formation_section, classifier_formations_experiences
import pytest


def test_docx_tables_are_read_in_document_order():
    document = Document()
    document.add_paragraph("AVANT")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "CELLULE"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "IMBRIQUE"
    document.add_paragraph("APRES")
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    text = lire_cv_docx(buffer)
    assert text.index("AVANT") < text.index("CELLULE") < text.index("IMBRIQUE") < text.index("APRES")


@pytest.mark.parametrize("heading", ["EXPERIENCES", "EXPÉRIENCES", "Expériences professionnelles"])
def test_formation_section_stops_before_experience(heading):
    text = f"FORMATIONS\n2020 - 2022 | Université Exemple | Master Informatique\n{heading}\n2022 - 2024 | Entreprise Démonstration | Développeur Python\n"
    formations = parse_formation_section(text)
    assert len(formations) == 1
    assert "Entreprise" not in formations[0]["etablissement"]


def test_ner_organizations_feed_experience_fallback():
    _, experiences = classifier_formations_experiences("Capgemini\n2020 - 2022", {"organisations": ["Capgemini"]}, [])
    assert [item["entreprise"] for item in experiences] == ["Capgemini"]


def test_regex_email_is_preserved():
    assert extraire_infos_cv("demo@example.invalid")["emails"] == ["demo@example.invalid"]
